# -*- coding: utf-8 -*-
"""
Chuyển toàn bộ file Excel (.xlsx) trong thư mục dữ liệu thành các file JSON
để trang web tĩnh tra cứu.

    Mỗi file Excel  ->  1 file JSON  (data/files/<thư mục con>/<id>.json)
    Danh sách file  ->  data/index.json  (web đọc file này để liệt kê)

Cách dùng:
    python tools/build_data.py                       # dùng thư mục mặc định
    python tools/build_data.py --src "D:\\DuLieu"     # chỉ định thư mục nguồn
    python tools/build_data.py --force               # build lại toàn bộ

Chỉ những file Excel mới / có thay đổi mới được đọc lại (dựa vào kích thước
và thời gian sửa file), nên lần chạy sau sẽ rất nhanh.
"""
import argparse
import datetime as dt
import json
import os
import re
import sys
import unicodedata
import warnings

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("Thiếu thư viện openpyxl. Cài bằng lệnh:  pip install openpyxl")

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

DEFAULT_SRC = r"E:\Work_CCE\CCE_SOHOA\CCE_dữ liệu cuối_17102026"
ROOT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_OUT = os.path.join(ROOT_DIR, "data")

# Tên cột chuẩn – dùng để gộp các biến thể/lỗi chính tả giữa các file.
CANONICAL_COLUMNS = [
    "STT", "Số hiệu bằng", "Số vào sổ gốc cấp văn bằng", "Tên văn bằng",
    "Ngành đào tạo", "Mã ngành đào tạo", "Chuyên ngành đào tạo",
    "Họ, chữ đệm và tên", "Mã người học", "Số định danh cá nhân", "Hộ chiếu",
    "Ngày, tháng, năm sinh", "Giới tính", "Dân tộc", "Quốc tịch", "Nơi sinh",
    "Tên trường đã học", "Niên khóa", "Xếp loại tốt nghiệp", "Trình độ đào tạo",
    "Bậc trình độ", "Hình thức đào tạo", "Ngày nhập học", "Thời gian đào tạo",
    "Tổng số tín chỉ", "Ngôn ngữ đào tạo", "Số quyết định thành lập hội đồng",
    "Địa danh nơi cơ quan cấp bằng đặt trụ sở", "Ngày tháng năm cấp bằng",
    "Tên cơ quan cấp bằng", "Chức danh người ký bằng",
    "Họ, chữ đệm, tên người ký bằng", "Tình trạng văn bằng", "Ghi chú",
    "Trạng thái số hóa",
    "Thông tin chưa bảo đảm đầy đủ, chính xác và nguyên nhân",
]
COLUMN_ALIASES = {
    "xep loai/hang tot nghiep": "Xếp loại tốt nghiệp",
    "chuyen nganh dao tao": "Chuyên ngành đào tạo",   # "đạo tạo", "đạo tào"
    "chuyen nganh dao tao.": "Chuyên ngành đào tạo",
    "trinh do dai hoc": "Trình độ đào tạo",
}

HEADER_HINTS = ("ho, chu dem va ten", "so hieu bang", "stt")


def strip_accents(text):
    text = unicodedata.normalize("NFD", str(text))
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    return text.replace("đ", "d").replace("Đ", "D")


def norm_key(text):
    return re.sub(r"\s+", " ", strip_accents(text).lower()).strip()


_CANON_BY_KEY = {norm_key(c): c for c in CANONICAL_COLUMNS}


def canonical_column(name):
    key = norm_key(name)
    if key in _CANON_BY_KEY:
        return _CANON_BY_KEY[key]
    if key in COLUMN_ALIASES:
        return COLUMN_ALIASES[key]
    # "đạo tào" -> "dao tao" sau khi bỏ dấu nên đã khớp ở trên; giữ nguyên nếu lạ.
    return re.sub(r"\s+", " ", str(name)).strip()


def slugify(text):
    text = strip_accents(text).lower()
    text = re.sub(r"[^a-z0-9]+", "-", text).strip("-")
    return text[:80] or "file"


def fmt_value(v):
    """Chuẩn hoá giá trị ô thành chuỗi."""
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        if v.hour == v.minute == v.second == 0:
            return v.strftime("%d/%m/%Y")
        return v.strftime("%d/%m/%Y %H:%M")
    if isinstance(v, dt.date):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, bool):
        return "Có" if v else "Không"
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else repr(v)
    return re.sub(r"[ \t]+", " ", str(v)).strip()


def excel_serial_to_date(v):
    """Ô ngày tháng lưu dạng số (General) -> date. Năm đơn lẻ (vd. 1978) giữ nguyên."""
    if isinstance(v, (int, float)) and not isinstance(v, bool) and 3000 <= v <= 80000:
        return dt.date(1899, 12, 30) + dt.timedelta(days=int(v))
    if isinstance(v, str) and re.fullmatch(r"\d{5}", v.strip()):
        return excel_serial_to_date(int(v.strip()))
    return v


def find_header(rows):
    """Trả về (vị trí dòng tiêu đề, tiêu đề bảng) trong 20 dòng đầu."""
    best = None
    for i, row in enumerate(rows[:20]):
        cells = [norm_key(c) for c in row if c not in (None, "")]
        if len(cells) < 3:
            continue
        score = sum(1 for h in HEADER_HINTS if h in cells)
        if score and (best is None or score > best[1]):
            best = (i, score)
    if best is None:
        return None, ""
    idx = best[0]
    title = ""
    for row in rows[:idx]:
        vals = [fmt_value(c) for c in row if c not in (None, "")]
        if vals:
            title = vals[0]
    return idx, title


def read_workbook(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        order = (["Data"] if "Data" in wb.sheetnames else []) + \
                [s for s in wb.sheetnames if s not in ("Data", "HuongDan", "DanhMuc")]
        for sheet_name in order:
            rows = [list(r) for r in wb[sheet_name].iter_rows(values_only=True)]
            hidx, title = find_header(rows)
            if hidx is None:
                continue
            header = [fmt_value(c) for c in rows[hidx]]
            while header and header[-1] == "":
                header.pop()
            columns, seen = [], set()
            for i, h in enumerate(header):
                name = canonical_column(h) if h else f"Cột {i + 1}"
                base, n = name, 2
                while name in seen:
                    name = f"{base} ({n})"
                    n += 1
                seen.add(name)
                columns.append(name)

            records = []
            date_cols = {i for i, c in enumerate(columns) if norm_key(c).startswith("ngay")}
            for row in rows[hidx + 1:]:
                vals = []
                for i in range(len(columns)):
                    raw = row[i] if i < len(row) else None
                    if i in date_cols:
                        raw = excel_serial_to_date(raw)
                    vals.append(fmt_value(raw))
                if not any(v for c, v in zip(columns, vals) if c != "STT"):
                    continue
                records.append(vals)

            # Bỏ các cột "Cột N" (không có tiêu đề) hoàn toàn trống
            drop = [i for i, c in enumerate(columns)
                    if c.startswith("Cột ") and not any(r[i] for r in records)]
            if drop:
                columns = [c for i, c in enumerate(columns) if i not in drop]
                records = [[v for i, v in enumerate(r) if i not in drop] for r in records]
            # Cắt bớt các ô trống ở cuối mỗi dòng cho gọn file JSON
            for r in records:
                while r and r[-1] == "":
                    r.pop()
            return {"sheet": sheet_name, "title": title,
                    "columns": columns, "records": records}
        raise ValueError("Không tìm thấy dòng tiêu đề (STT / Họ, chữ đệm và tên ...)")
    finally:
        wb.close()


def write_json_and_js(json_path, obj, js_prefix=None, pretty=False):
    """Ghi cả file .json và file .js để vừa có json vừa chạy được offline bấm đúp index.html."""
    os.makedirs(os.path.dirname(json_path), exist_ok=True)
    tmp_json = json_path + ".tmp"
    with open(tmp_json, "w", encoding="utf-8") as f:
        if pretty:
            json.dump(obj, f, ensure_ascii=False, indent=2)
        else:
            json.dump(obj, f, ensure_ascii=False, separators=(",", ":"))
    os.replace(tmp_json, json_path)

    if js_prefix:
        js_path = os.path.splitext(json_path)[0] + ".js"
        tmp_js = js_path + ".tmp"
        with open(tmp_js, "w", encoding="utf-8") as f:
            f.write(js_prefix + "(" if js_prefix.endswith("__LOAD_FILE__") else js_prefix)
            if pretty:
                json.dump(obj, f, ensure_ascii=False, indent=2)
            else:
                json.dump(obj, f, ensure_ascii=False, separators=(",", ":"))
            f.write(");\n" if js_prefix.endswith("__LOAD_FILE__") else ";\n")
        os.replace(tmp_js, js_path)


def write_json(path, obj, pretty=False):
    write_json_and_js(path, obj, pretty=pretty)


def main():
    ap = argparse.ArgumentParser(description="Chuyển Excel -> JSON cho web tra cứu")
    ap.add_argument("--src", default=DEFAULT_SRC, help="Thư mục chứa file Excel")
    ap.add_argument("--out", default=DEFAULT_OUT, help="Thư mục xuất dữ liệu (mặc định ./data)")
    ap.add_argument("--force", action="store_true", help="Đọc lại toàn bộ file Excel")
    args = ap.parse_args()

    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

    src = os.path.abspath(args.src)
    out = os.path.abspath(args.out)
    if not os.path.isdir(src):
        sys.exit(f"Không tìm thấy thư mục nguồn: {src}")

    index_path = os.path.join(out, "index.json")
    old = {}
    if os.path.isfile(index_path) and not args.force:
        try:
            with open(index_path, encoding="utf-8") as f:
                old = {e["source"]: e for e in json.load(f).get("files", [])}
        except Exception:
            old = {}

    excel_files = []
    for dirpath, dirnames, filenames in os.walk(src):
        dirnames.sort()
        for fn in sorted(filenames):
            if fn.lower().endswith((".xlsx", ".xlsm")) and not fn.startswith("~$"):
                excel_files.append(os.path.join(dirpath, fn))

    entries, used_ids, errors = [], set(), []
    reused = parsed = 0
    total = len(excel_files)
    for n, path in enumerate(excel_files, 1):
        rel = os.path.relpath(path, src).replace("\\", "/")
        group = rel.split("/")[0] if "/" in rel else "Khác"
        name = os.path.splitext(os.path.basename(path))[0]
        st = os.stat(path)

        base_id = f"{slugify(group)}-{slugify(name)}"
        fid, k = base_id, 2
        while fid in used_ids:
            fid, k = f"{base_id}-{k}", k + 1
        used_ids.add(fid)
        json_rel = f"files/{slugify(group)}/{fid}.json"
        json_abs = os.path.join(out, json_rel)

        prev = old.get(rel)
        if (prev and prev.get("size") == st.st_size and prev.get("mtime") == int(st.st_mtime)
                and prev.get("path") == json_rel and os.path.isfile(json_abs)):
            entries.append(prev)
            reused += 1
            continue

        print(f"[{n}/{total}] {rel}")
        try:
            data = read_workbook(path)
        except Exception as e:  # noqa: BLE001
            errors.append((rel, str(e)))
            print(f"   !! Lỗi: {e}")
            continue

        doc = {
            "id": fid,
            "name": name,
            "fileName": os.path.basename(path),
            "group": group,
            "source": rel,
            "sheet": data["sheet"],
            "title": data["title"],
            "columns": data["columns"],
            "count": len(data["records"]),
            # Mỗi dòng là mảng giá trị theo đúng thứ tự "columns"
            "rows": data["records"],
        }
        write_json_and_js(json_abs, doc, js_prefix="window.__LOAD_FILE__")
        entry = {
            "id": fid, "name": name, "fileName": os.path.basename(path),
            "group": group, "source": rel, "path": json_rel,
            "title": data["title"], "count": len(data["records"]),
            "size": st.st_size, "mtime": int(st.st_mtime),
        }
        extra = [c for c in data["columns"] if c not in CANONICAL_COLUMNS]
        if extra:
            entry["extraColumns"] = extra
        entries.append(entry)
        parsed += 1

    # Xoá JSON mồ côi (file Excel đã bị xoá / đổi tên)
    keep = {os.path.normcase(os.path.join(out, e["path"])) for e in entries}
    files_dir = os.path.join(out, "files")
    removed = 0
    for dirpath, _, filenames in os.walk(files_dir):
        for fn in filenames:
            p = os.path.join(dirpath, fn)
            if fn.endswith(".json") and os.path.normcase(p) not in keep:
                os.remove(p)
                removed += 1

    all_cols = list(CANONICAL_COLUMNS)
    for e in entries:
        for c in e.get("extraColumns", []):
            if c not in all_cols:
                all_cols.append(c)
    index = {
        "generatedAt": dt.datetime.now().strftime("%d/%m/%Y %H:%M"),
        "totalFiles": len(entries),
        "totalRecords": sum(e["count"] for e in entries),
        "columns": all_cols,
        "groups": sorted({e["group"] for e in entries}),
        "files": entries,
    }
    write_json_and_js(index_path, index, js_prefix="window.__DATA_MANIFEST__ = ", pretty=True)

    # Xuất danh sách tên người học phục vụ gợi ý tức thì (autocomplete)
    names_set = set()
    for e in entries:
        json_abs = os.path.join(out, e["path"])
        if os.path.isfile(json_abs):
            try:
                with open(json_abs, "r", encoding="utf-8") as f:
                    doc = json.load(f)
                    name_idx = next((i for i, c in enumerate(doc.get("columns", [])) if "họ" in c.lower()), -1)
                    if name_idx != -1:
                        for r in doc.get("rows", []):
                            if name_idx < len(r):
                                nm = r[name_idx].strip()
                                if nm and len(nm) >= 2:
                                    names_set.add(nm)
            except Exception:
                pass
    names_list = sorted(list(names_set))
    names_json = os.path.join(out, "names.json")
    write_json_and_js(names_json, names_list, js_prefix="window.__NAMES__ = ")
    print(f"Danh sách tên gợi ý: {len(names_list)} tên")

    # Tạo 1 file gộp duy nhất (bundle) chứa toàn bộ dữ liệu 620 file
    # Giúp nạp tức thì 1 lần duy nhất (~0.1s), F5 là chạy ngay, không gửi 620 request lẻ tẻ
    bundle_docs = []
    for e in entries:
        json_abs = os.path.join(out, e["path"])
        if os.path.isfile(json_abs):
            try:
                with open(json_abs, "r", encoding="utf-8") as f:
                    d = json.load(f)
                    bundle_docs.append({
                        "id": d["id"],
                        "fileName": d["fileName"],
                        "group": d["group"],
                        "title": d.get("title", ""),
                        "columns": d["columns"],
                        "rows": d["rows"],
                    })
            except Exception:
                pass

    bundle_path = os.path.join(out, "bundle.json")
    write_json_and_js(bundle_path, bundle_docs, js_prefix="window.__DATA_BUNDLE__ = ")
    print(f"Đã tạo file bundle tổng hợp ({len(bundle_docs)} file)")

    print("-" * 60)
    print(f"Tổng số file Excel : {total}")
    print(f"Đọc mới / cập nhật : {parsed}")
    print(f"Giữ nguyên (cache) : {reused}")
    print(f"JSON cũ đã xoá     : {removed}")
    print(f"Tổng số bản ghi    : {index['totalRecords']}")
    print(f"Đã ghi             : {index_path}")
    if errors:
        print(f"\n{len(errors)} file lỗi:")
        for rel, msg in errors:
            print(f"  - {rel}: {msg}")


if __name__ == "__main__":
    main()
